"""A Sua Drogaria — Plataforma de delivery (FastAPI + MongoDB).

Loja do cliente + painel administrativo compartilham o mesmo banco.
"""
import csv
import io
import logging
import os
import re
import unicodedata
import uuid
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Annotated, Any, Dict, List, Optional

import jwt
import openpyxl
import requests
from bson import ObjectId
from dotenv import load_dotenv
from fastapi import APIRouter, Body, Depends, FastAPI, File, Header, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from pydantic import BaseModel, BeforeValidator, ConfigDict, Field
from starlette.concurrency import run_in_threadpool

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

mongo_url = os.environ["MONGO_URL"]
from motor.motor_asyncio import AsyncIOMotorClient  # noqa: E402

client = AsyncIOMotorClient(mongo_url)
db = client[os.environ["DB_NAME"]]

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(name)s - %(levelname)s - %(message)s")
logger = logging.getLogger("asua-drogaria")

# ---------------------------------------------------------------------------
# Configurações da loja (via .env)
# ---------------------------------------------------------------------------
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "")
JWT_SECRET = os.environ.get("JWT_SECRET", "troque-esta-chave")
PIX_KEY = os.environ.get("PIX_KEY", "")
PIX_NAME = os.environ.get("PIX_NAME", "A Sua Drogaria")
STORE_NAME = os.environ.get("STORE_NAME", "A Sua Drogaria")
STORE_PHONE = os.environ.get("STORE_PHONE", "21 97921-6950")
STORE_CITY = os.environ.get("STORE_CITY", "Duque de Caxias, RJ")
DELIVERY_FEE = float(os.environ.get("DELIVERY_FEE", "0"))

VALID_STATUSES = ["pendente", "em_separacao", "saiu_entrega", "finalizado", "cancelado"]
VALID_PAYMENTS = ["pix", "dinheiro", "cartao"]

# ---------------------------------------------------------------------------
# Object Storage (Emergent) — upload/foto de produtos
# ---------------------------------------------------------------------------
STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
EMERGENT_KEY = os.environ.get("EMERGENT_LLM_KEY")
APP_NAME = "asua-drogaria"
storage_key = None


def init_storage():
    global storage_key
    if storage_key:
        return storage_key
    resp = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_KEY}, timeout=30)
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key


def put_object(path: str, data: bytes, content_type: str) -> dict:
    key = init_storage()
    resp = requests.put(
        f"{STORAGE_URL}/objects/{path}",
        headers={"X-Storage-Key": key, "Content-Type": content_type},
        data=data,
        timeout=120,
    )
    if resp.status_code == 503:  # storage_key expirado → re-init uma vez
        global storage_key
        storage_key = None
        key = init_storage()
        resp = requests.put(
            f"{STORAGE_URL}/objects/{path}",
            headers={"X-Storage-Key": key, "Content-Type": content_type},
            data=data,
            timeout=120,
        )
    resp.raise_for_status()
    return resp.json()


def get_object(path: str):
    key = init_storage()
    resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    if resp.status_code == 503:
        global storage_key
        storage_key = None
        key = init_storage()
        resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    resp.raise_for_status()
    return resp.content, resp.headers.get("Content-Type", "application/octet-stream")


# ---------------------------------------------------------------------------
# Modelos (MongoDB adherence)
# ---------------------------------------------------------------------------
PyObjectId = Annotated[str, BeforeValidator(lambda v: str(v))]


class BaseDocument(BaseModel):
    model_config = ConfigDict(populate_by_name=True)
    id: Optional[PyObjectId] = Field(None, alias="_id")

    def to_mongo(self) -> dict:
        data = self.model_dump(by_alias=True)
        if data.get("_id") is None:
            data.pop("_id", None)
        return data

    @classmethod
    def from_mongo(cls, doc: dict):
        return cls(**doc)


class Product(BaseDocument):
    name: str
    description: str = ""
    category: str = "medicamentos"
    price: float
    promo_price: Optional[float] = None
    stock: int = 0
    barcode: Optional[str] = None
    external_id: Optional[str] = None
    image_url: Optional[str] = None
    active: bool = True
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))
    updated_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class ProductCreate(BaseModel):
    name: str = Field(min_length=2)
    description: str = ""
    category: str = "medicamentos"
    price: float = Field(ge=0)
    promo_price: Optional[float] = None
    stock: int = Field(default=0, ge=0)
    barcode: Optional[str] = None
    image_url: Optional[str] = None


class ProductUpdate(BaseModel):
    name: Optional[str] = None
    description: Optional[str] = None
    category: Optional[str] = None
    price: Optional[float] = None
    promo_price: Optional[float] = None
    stock: Optional[int] = None
    barcode: Optional[str] = None
    image_url: Optional[str] = None
    active: Optional[bool] = None


class OrderItemIn(BaseModel):
    product_id: str
    qty: int = Field(ge=1, le=99)


class OrderItem(BaseModel):
    product_id: str
    name: str
    price: float
    qty: int


class OrderCreate(BaseModel):
    customer_name: str = Field(min_length=3)
    customer_phone: str = Field(min_length=8)
    address: str = Field(min_length=5)
    neighborhood: str = Field(min_length=2)
    complement: str = ""
    payment_method: str
    items: List[OrderItemIn] = Field(min_length=1)


class Order(BaseDocument):
    code: str
    customer_name: str
    customer_phone: str
    address: str
    neighborhood: str
    complement: str = ""
    payment_method: str
    items: List[OrderItem] = []
    subtotal: float
    delivery_fee: float = 0.0
    total: float
    status: str = "pendente"
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))
    updated_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class StatusUpdate(BaseModel):
    status: str


class AdminLogin(BaseModel):
    password: str


# ---------------------------------------------------------------------------
# Autenticação do admin (senha única → JWT)
# ---------------------------------------------------------------------------
def require_admin(authorization: Optional[str] = Header(None)):
    if not ADMIN_PASSWORD:
        raise HTTPException(status_code=500, detail="ADMIN_PASSWORD não configurado no servidor")
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Faça login para acessar o painel")
    token = authorization.split(" ", 1)[1]
    try:
        jwt.decode(token, JWT_SECRET, algorithms=["HS256"])
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Sessão expirada. Entre novamente.")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Sessão inválida. Entre novamente.")


# ---------------------------------------------------------------------------
# Utilidades de importação (planilha / webhook)
# ---------------------------------------------------------------------------
def _norm(s: Any) -> str:
    s = unicodedata.normalize("NFKD", str(s or "").strip().lower())
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[\s_\-./()]+", " ", s).strip()


FIELD_ALIASES: Dict[str, set] = {
    "barcode": {"codigo de barras", "codigo barras", "cod de barras", "cod barras", "ean", "gtin", "barcode", "codigo", "cod"},
    "external_id": {"id", "identificador", "codigo interno", "sku", "referencia", "ref"},
    "name": {"nome", "produto", "name", "medicamento", "descricao produto", "titulo", "nome do produto"},
    "price": {"preco", "valor", "price", "preco venda", "preco unitario", "preco de venda", "pvp", "preco atual"},
    "promo_price": {"preco promocional", "preco promo", "promo", "promocao", "preco promocional se houver", "preco especial"},
    "stock": {"estoque", "qtd", "qtde", "quantidade", "quantidade estoque", "qty", "stock", "saldo", "qtd estoque"},
    "category": {"categoria", "secao", "departamento", "categoria produto", "grupo"},
    "description": {"descricao", "detalhes", "description", "observacoes", "complemento descricao"},
    "image_url": {"imagem", "foto", "image", "url imagem", "link imagem", "imagem url"},
}

CATEGORY_MAP: Dict[str, set] = {
    "medicamentos": {"medicamento", "medicamentos", "remedio", "remedios", "farmacia", "medicamento e saude"},
    "higiene": {"higiene", "higiene e beleza", "higiene pessoal", "cuidados pessoais", "higiene e cuidados"},
    "cosmeticos": {"cosmetico", "cosmeticos", "beleza", "perfumaria", "dermocosmetico", "cosmeticos e beleza"},
    "suplementos": {"suplemento", "suplementos", "vitaminas", "nutricao", "esportivo", "suplementos e vitaminas"},
}


def normalize_category(value: Any) -> Optional[str]:
    if value in (None, ""):
        return None
    n = _norm(value)
    for cat, aliases in CATEGORY_MAP.items():
        if n in aliases:
            return cat
    for cat in CATEGORY_MAP:
        if cat in n:
            return cat
    return None


def parse_money(value: Any) -> Optional[float]:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    s = re.sub(r"[^\d,.\-]", "", str(value))
    if not s:
        return None
    if "," in s and s.rfind(",") > s.rfind("."):
        s = s.replace(".", "").replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    try:
        return round(float(s), 2)
    except ValueError:
        return None


def parse_int(value: Any) -> Optional[int]:
    if value in (None, ""):
        return None
    try:
        return int(float(str(value).replace(",", ".")))
    except ValueError:
        return None


def map_record(record: dict) -> dict:
    mapped: Dict[str, Any] = {}
    for key, value in record.items():
        n = _norm(key)
        if value in (None, ""):
            continue
        for field, aliases in FIELD_ALIASES.items():
            if n in aliases and field not in mapped:
                mapped[field] = value
    return mapped


def rows_from_csv(data: bytes) -> List[dict]:
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    delim = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delim)
    return [dict(r) for r in reader if any((v or "").strip() if isinstance(v, str) else v is not None for v in r.values())]


def rows_from_xlsx(data: bytes) -> List[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for row in rows[1:]:
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        out.append({headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]})
    return out


async def upsert_product(rec: dict) -> str:
    """Atualiza/cria produto por código de barras (ou id externo). Retorna status."""
    barcode = str(rec.get("barcode")).strip() if rec.get("barcode") not in (None, "") else None
    external_id = str(rec.get("external_id")).strip() if rec.get("external_id") not in (None, "") else None
    price = parse_money(rec.get("price"))
    promo = parse_money(rec.get("promo_price"))
    stock = parse_int(rec.get("stock"))
    category = normalize_category(rec.get("category"))
    name = str(rec.get("name")).strip() if rec.get("name") else None

    if not barcode and not external_id and not name:
        return "skipped"

    query: dict = {}
    if barcode:
        query["barcode"] = barcode
    elif external_id:
        query["external_id"] = external_id
    else:
        query["name"] = name

    existing = await db.products.find_one(query)
    now = datetime.now(timezone.utc)

    updates: dict = {}
    if price is not None:
        updates["price"] = price
    if promo is not None:
        updates["promo_price"] = promo
    if stock is not None:
        updates["stock"] = max(0, stock)
    if category:
        updates["category"] = category
    if name:
        updates["name"] = name
    if rec.get("description"):
        updates["description"] = str(rec["description"]).strip()
    if rec.get("image_url"):
        updates["image_url"] = str(rec["image_url"]).strip()
    if external_id:
        updates["external_id"] = external_id

    if existing:
        if not updates:
            return "skipped"
        await db.products.update_one({"_id": existing["_id"]}, {"$set": {**updates, "updated_at": now}})
        return "updated"

    if not name:
        return "skipped"
    doc = Product(
        name=name,
        description=str(rec.get("description") or "").strip(),
        category=category or "medicamentos",
        price=price if price is not None else 0.0,
        promo_price=promo,
        stock=stock if stock is not None else 0,
        barcode=barcode,
        external_id=external_id,
        image_url=str(rec["image_url"]).strip() if rec.get("image_url") else None,
        created_at=now,
        updated_at=now,
    ).to_mongo()
    await db.products.insert_one(doc)
    return "created"


# ---------------------------------------------------------------------------
# Seed — produtos de demonstração
# ---------------------------------------------------------------------------
DEMO_PRODUCTS = [
    # Medicamentos
    ("Dipirona 500mg 20 comprimidos", "medicamentos", "Analgésico e antitérmico. Uso adulto e pediátrico conforme orientação médica.", 9.90, 7.90, 42, "7891058001231", "https://images.unsplash.com/photo-1550572017-edd951b55104?w=600&q=80&auto=format&fit=crop"),
    ("Paracetamol 750mg 20 comprimidos", "medicamentos", "Alívio da dor e da febre. Não exceder a dose recomendada.", 12.50, None, 35, "7891058004569", "https://images.unsplash.com/photo-1587854692152-cbe660dbde88?w=600&q=80&auto=format&fit=crop"),
    ("Ibuprofeno 600mg 30 cápsulas", "medicamentos", "Anti-inflamatório não esteroide para dores e inflamações.", 24.90, 19.90, 28, "7891058007892", "https://images.unsplash.com/photo-1550572017-edd951b55104?w=600&q=80&auto=format&fit=crop"),
    ("Omeprazol 20mg 28 cápsulas", "medicamentos", "Reduz a acidez do estômago. Uso contínuo sob orientação médica.", 32.90, None, 18, "7891058003211", "https://images.unsplash.com/photo-1471864190281-a93a3070b6de?w=600&q=80&auto=format&fit=crop"),
    ("Loratadina 10mg 12 comprimidos", "medicamentos", "Antialérgico para rinite, alergias e urticária.", 14.90, None, 22, "7891058006548", "https://images.unsplash.com/photo-1471864190281-a93a3070b6de?w=600&q=80&auto=format&fit=crop"),
    ("Soro Fisiológico 0,9% 500ml", "medicamentos", "Solução fisiológica para uso nasal e limpeza de ferimentos.", 8.50, None, 30, "7891058009876", "https://images.unsplash.com/photo-1584362917165-526a968579e8?w=600&q=80&auto=format&fit=crop"),
    ("Álcool Gel 70% 500ml", "medicamentos", "Antisséptico para higienização das mãos e superfícies.", 12.90, 9.90, 50, "7891058001354", "https://images.unsplash.com/photo-1603398938378-e54eab446dde?w=600&q=80&auto=format&fit=crop"),
    # Higiene
    ("Creme Dental Anticáries 90g", "higiene", "Proteção contra cáries com flúor. Limpeza completa.", 7.90, None, 60, "7891024001111", "https://images.unsplash.com/photo-1607613009820-a29f7bb81c04?w=600&q=80&auto=format&fit=crop"),
    ("Shampoo Hidratante 350ml", "higiene", "Limpeza suave com hidratação prolongada para todos os tipos de cabelo.", 22.90, None, 25, "7891024002222", "https://images.unsplash.com/photo-1626806787461-102c1bfaaea1?w=600&q=80&auto=format&fit=crop"),
    ("Sabonete Líquido 250ml", "higiene", "Sabonete neutro dermatologicamente testado para uso diário.", 11.90, 9.50, 40, "7891024003333", "https://images.unsplash.com/photo-1556228720-195a672e8a03?w=600&q=80&auto=format&fit=crop"),
    ("Desodorante Aerosol 150ml", "higiene", "Proteção antitranspirante por 48 horas.", 18.90, 15.90, 33, "7891024004444", "https://images.unsplash.com/photo-1556228578-8c89e6adf883?w=600&q=80&auto=format&fit=crop"),
    ("Papel Higiênico Folha Dupla 12 rolos", "higiene", "Macio e absorvente, folha dupla.", 21.90, None, 20, "7891024005555", "https://images.unsplash.com/photo-1585435557343-3b092031a831?w=600&q=80&auto=format&fit=crop"),
    # Cosméticos
    ("Protetor Solar FPS 60 120ml", "cosmeticos", "Proteção UVA/UVB, toque seco, ideal para o dia a dia.", 59.90, 49.90, 15, "7891024006666", "https://images.unsplash.com/photo-1620916566398-39f1143ab7be?w=600&q=80&auto=format&fit=crop"),
    ("Hidratante Corporal 400ml", "cosmeticos", "Hidratação profunda com vitamina E para pele seca.", 34.90, None, 17, "7891024007777", "https://images.unsplash.com/photo-1571781926291-c477ebfd024b?w=600&q=80&auto=format&fit=crop"),
    ("Perfume Feminino Flor de Lótus 75ml", "cosmeticos", "Fragrância floral marcante com fixação de longa duração.", 129.90, 99.90, 8, "7891024008888", "https://images.unsplash.com/photo-1598300042247-d088f8ab3a91?w=600&q=80&auto=format&fit=crop"),
    ("Kit de Maquiagem Essencial", "cosmeticos", "Kit com sombras, batom e pincéis para o dia a dia.", 49.90, None, 12, "7891024009999", "https://images.unsplash.com/photo-1571875257727-256c39da42af?w=600&q=80&auto=format&fit=crop"),
    # Suplementos
    ("Whey Protein Baunilha 900g", "suplementos", "Proteína de alto valor biológico para ganho de massa muscular.", 119.90, 99.90, 10, "7891034001111", "https://images.unsplash.com/photo-1607619056574-7b8d3ee536b2?w=600&q=80&auto=format&fit=crop"),
    ("Vitamina C 1000mg 60 cápsulas", "suplementos", "Auxilia o sistema imunológico e a saúde da pele.", 39.90, None, 26, "7891034002222", "https://images.unsplash.com/photo-1585435557343-3b092031a831?w=600&q=80&auto=format&fit=crop"),
    ("Creatina Monohidratada 150g", "suplementos", "Pura, para força e desempenho nos treinos.", 69.90, None, 14, "7891034003333", "https://images.unsplash.com/photo-1607619056574-7b8d3ee536b2?w=600&q=80&auto=format&fit=crop"),
    ("Ômega 3 120 cápsulas", "suplementos", "Ácidos graxos essenciais para a saúde do coração.", 54.90, 44.90, 19, "7891034004444", "https://images.unsplash.com/photo-1471864190281-a93a3070b6de?w=600&q=80&auto=format&fit=crop"),
]


async def seed_products_if_empty():
    if await db.products.count_documents({}) > 0:
        return
    docs = []
    for name, cat, desc, price, promo, stock, barcode, image in DEMO_PRODUCTS:
        docs.append(Product(
            name=name, description=desc, category=cat, price=price, promo_price=promo,
            stock=stock, barcode=barcode, image_url=image,
        ).to_mongo())
    await db.products.insert_many(docs)
    logger.info("Seed: %d produtos de demonstração inseridos", len(docs))


@asynccontextmanager
async def lifespan(app: FastAPI):
    try:
        await run_in_threadpool(init_storage)
        logger.info("Object Storage inicializado")
    except Exception as e:
        logger.error("Falha ao inicializar Object Storage: %s", e)
    try:
        await seed_products_if_empty()
    except Exception as e:
        logger.error("Falha no seed de produtos: %s", e)
    yield
    client.close()


app = FastAPI(title="A Sua Drogaria API", lifespan=lifespan)
api_router = APIRouter(prefix="/api")


@api_router.post("/admin/login")
async def admin_login(body: AdminLogin):
    if not ADMIN_PASSWORD or body.password != ADMIN_PASSWORD:
        raise HTTPException(status_code=401, detail="Senha incorreta. Tente novamente.")
    expires = datetime.now(timezone.utc) + timedelta(days=30)
    token = jwt.encode({"role": "admin", "exp": expires}, JWT_SECRET, algorithm="HS256")
    return {"token": token}


@api_router.get("/")
async def root():
    return {
        "app": STORE_NAME,
        "status": "ok",
        "webhook": {
            "url": "/api/webhook/products",
            "metodo": "POST",
            "exemplo": [
                {"id": "TR-001", "codigo_barras": "7891058001231", "preco": 9.90, "estoque": 42},
            ],
            "campos_aceitos": "codigo_barras/barcode/ean, preco/price, estoque/stock, nome, categoria, id",
        },
    }


@api_router.get("/config")
async def get_config():
    return {
        "name": STORE_NAME,
        "phone": STORE_PHONE,
        "city": STORE_CITY,
        "pix_key": PIX_KEY,
        "pix_name": PIX_NAME,
        "delivery_fee": DELIVERY_FEE,
        "delivery_time": "30 a 60 minutos",
    }


# ---------------------------------------------------------------------------
# Produtos (público — loja do cliente)
# ---------------------------------------------------------------------------
@api_router.get("/products")
async def list_products(category: Optional[str] = None, search: Optional[str] = None, limit: int = 300):
    query: dict = {"active": {"$ne": False}}
    if category and category != "todos":
        query["category"] = category
    if search:
        query["name"] = {"$regex": re.escape(search.strip()), "$options": "i"}
    docs = await db.products.find(query).sort("name", 1).to_list(limit)
    return [Product.from_mongo(d).model_dump() for d in docs]


@api_router.get("/products/{product_id}")
async def get_product(product_id: str):
    try:
        doc = await db.products.find_one({"_id": ObjectId(product_id)})
    except Exception:
        raise HTTPException(status_code=404, detail="Produto não encontrado")
    if not doc or doc.get("active") is False:
        raise HTTPException(status_code=404, detail="Produto não encontrado")
    return Product.from_mongo(doc).model_dump()


# ---------------------------------------------------------------------------
# Pedidos (público — cria pedido; leitura por código para a tela de confirmação)
# ---------------------------------------------------------------------------
@api_router.post("/orders")
async def create_order(payload: OrderCreate):
    if payload.payment_method not in VALID_PAYMENTS:
        raise HTTPException(status_code=400, detail="Forma de pagamento inválida")

    items: List[OrderItem] = []
    subtotal = 0.0
    now = datetime.now(timezone.utc)

    for item in payload.items:
        try:
            doc = await db.products.find_one({"_id": ObjectId(item.product_id)})
        except Exception:
            raise HTTPException(status_code=400, detail="Produto inválido no carrinho")
        if not doc or doc.get("active") is False:
            raise HTTPException(status_code=400, detail="Produto indisponível no carrinho")
        product = Product.from_mongo(doc)
        unit_price = product.promo_price if (product.promo_price and product.promo_price < product.price) else product.price
        if product.stock < item.qty:
            raise HTTPException(status_code=409, detail=f"Estoque insuficiente: {product.name}")
        items.append(OrderItem(product_id=item.product_id, name=product.name, price=unit_price, qty=item.qty))
        subtotal += unit_price * item.qty

    # Baixa o estoque de forma segura (race-aware)
    for item in items:
        result = await db.products.update_one(
            {"_id": ObjectId(item.product_id), "stock": {"$gte": item.qty}},
            {"$inc": {"stock": -item.qty}, "$set": {"updated_at": now}},
        )
        if result.matched_count == 0:
            raise HTTPException(status_code=409, detail=f"Estoque insuficiente: {item.name}")

    counter = await db.counters.find_one_and_update(
        {"_id": "orders"}, {"$inc": {"seq": 1}}, upsert=True, return_document=True
    )
    code = str(1000 + counter["seq"])

    order = Order(
        code=code,
        customer_name=payload.customer_name.strip(),
        customer_phone=payload.customer_phone.strip(),
        address=payload.address.strip(),
        neighborhood=payload.neighborhood.strip(),
        complement=payload.complement.strip(),
        payment_method=payload.payment_method,
        items=items,
        subtotal=round(subtotal, 2),
        delivery_fee=DELIVERY_FEE,
        total=round(subtotal + DELIVERY_FEE, 2),
        status="pendente",
        created_at=now,
        updated_at=now,
    )
    result = await db.orders.insert_one(order.to_mongo())
    order.id = str(result.inserted_id)
    return order.model_dump()


def _mask_phone(phone: str) -> str:
    digits = re.sub(r"\D", "", phone)
    if len(digits) <= 4:
        return "****"
    return f"({digits[:2]}) {digits[2:-4]}-{digits[-4:]}"


@api_router.get("/orders/{code}")
async def get_order_by_code(code: str):
    doc = await db.orders.find_one({"code": code})
    if not doc:
        raise HTTPException(status_code=404, detail="Pedido não encontrado")
    order = Order.from_mongo(doc)
    data = order.model_dump()
    data["customer_phone"] = _mask_phone(data["customer_phone"])
    return data


# ---------------------------------------------------------------------------
# Admin — pedidos (Kanban em tempo real)
# ---------------------------------------------------------------------------
@api_router.get("/admin/orders")
async def admin_list_orders(_: None = Depends(require_admin)):
    docs = await db.orders.find().sort("created_at", -1).to_list(300)
    return [Order.from_mongo(d).model_dump() for d in docs]


@api_router.patch("/admin/orders/{order_id}/status")
async def admin_update_status(order_id: str, payload: StatusUpdate, _: None = Depends(require_admin)):
    if payload.status not in VALID_STATUSES:
        raise HTTPException(status_code=400, detail="Status inválido")
    
    try:
        order_doc = await db.orders.find_one({"_id": ObjectId(order_id)})
    except Exception:
        raise HTTPException(status_code=404, detail="Pedido não encontrado")
        
    if not order_doc:
        raise HTTPException(status_code=404, detail="Pedido não encontrado")
        
    old_status = order_doc.get("status")
    new_status = payload.status
    
    # Se o pedido foi cancelado agora e antes não estava cancelado, devolve o stock automaticamente
    if new_status == "cancelado" and old_status != "cancelado":
        for item in order_doc.get("items", []):
            await db.products.update_one(
                {"_id": ObjectId(item["product_id"])},
                {"$inc": {"stock": item["qty"]}, "$set": {"updated_at": datetime.now(timezone.utc)}}
            )

    result = await db.orders.update_one(
        {"_id": ObjectId(order_id)},
        {"$set": {"status": new_status, "updated_at": datetime.now(timezone.utc)}}
    )
    
    doc = await db.orders.find_one({"_id": ObjectId(order_id)})
    return Order.from_mongo(doc).model_dump()


# ---------------------------------------------------------------------------
# Admin — estoque (CRUD de produtos)
# ---------------------------------------------------------------------------
@api_router.get("/admin/products")
async def admin_list_products(_: None = Depends(require_admin)):
    docs = await db.products.find().sort("name", 1).to_list(1000)
    return [Product.from_mongo(d).model_dump() for d in docs]


@api_router.post("/admin/products")
async def admin_create_product(payload: ProductCreate, _: None = Depends(require_admin)):
    now = datetime.now(timezone.utc)
    product = Product(
        name=payload.name.strip(),
        description=payload.description.strip(),
        category=payload.category,
        price=payload.price,
        promo_price=payload.promo_price,
        stock=payload.stock,
        barcode=payload.barcode,
        image_url=payload.image_url,
        created_at=now,
        updated_at=now,
    )
    result = await db.products.insert_one(product.to_mongo())
    product.id = str(result.inserted_id)
    return product.model_dump()


@api_router.put("/admin/products/{product_id}")
async def admin_update_product(product_id: str, payload: ProductUpdate, _: None = Depends(require_admin)):
    updates: dict = {}
    for field in ("name", "description", "category", "price", "stock", "barcode", "image_url", "active"):
        if field in payload.model_fields_set and getattr(payload, field) is not None:
            updates[field] = getattr(payload, field)
    if "promo_price" in payload.model_fields_set:
        updates["promo_price"] = payload.promo_price
    if not updates:
        raise HTTPException(status_code=400, detail="Nada para atualizar")
    updates["updated_at"] = datetime.now(timezone.utc)
    try:
        result = await db.products.update_one({"_id": ObjectId(product_id)}, {"$set": updates})
    except Exception:
        raise HTTPException(status_code=404, detail="Produto não encontrado")
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Produto não encontrado")
    doc = await db.products.find_one({"_id": ObjectId(product_id)})
    return Product.from_mongo(doc).model_dump()


@api_router.delete("/admin/products/{product_id}")
async def admin_delete_product(product_id: str, _: None = Depends(require_admin)):
    # Soft delete — nunca destruir dados
    result = await db.products.update_one(
        {"_id": ObjectId(product_id)},
        {"$set": {"active": False, "deleted_at": datetime.now(timezone.utc), "updated_at": datetime.now(timezone.utc)}},
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Produto não encontrado")
    return {"ok": True}


# ---------------------------------------------------------------------------
# Webhook para sistemas de farmácia (Trier, InovaFarma, Linx Big...)
# ---------------------------------------------------------------------------
@api_router.post("/webhook/products")
async def webhook_products(payload: Any = Body(...)):
    records = payload
    if isinstance(records, dict):
        records = records.get("products") or records.get("data") or records.get("items") or [records]
    if not isinstance(records, list):
        raise HTTPException(status_code=400, detail="Envie um JSON array de produtos")

    counters = {"updated": 0, "created": 0, "skipped": 0}
    for raw in records:
        if not isinstance(raw, dict):
            counters["skipped"] += 1
            continue
        status = await upsert_product(map_record(raw))
        counters[status] += 1
    return {"ok": True, **counters, "total_recebido": len(records)}


# ---------------------------------------------------------------------------
# Admin — importação de planilha Excel/CSV
# ---------------------------------------------------------------------------
@api_router.post("/admin/import")
async def admin_import(file: UploadFile = File(...), _: None = Depends(require_admin)):
    data = await file.read()
    if not data:
        raise HTTPException(status_code=400, detail="Arquivo vazio")
    if len(data) > 10 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Arquivo muito grande (máx. 10 MB)")

    fname = (file.filename or "").lower()
    try:
        if fname.endswith((".xlsx", ".xlsm")):
            rows = rows_from_xlsx(data)
        elif fname.endswith((".csv", ".txt")):
            rows = rows_from_csv(data)
        else:
            try:
                rows = rows_from_csv(data)
            except Exception:
                raise HTTPException(status_code=400, detail="Formato não suportado. Envie .xlsx ou .csv")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="Não foi possível ler a planilha")

    if not rows:
        raise HTTPException(status_code=400, detail="Nenhuma linha encontrada na planilha")

    counters = {"updated": 0, "created": 0, "skipped": 0}
    errors: List[str] = []
    for i, raw in enumerate(rows, start=2):  # linha 1 = cabeçalho
        try:
            status = await upsert_product(map_record(raw))
            counters[status] += 1
        except Exception as e:
            errors.append(f"Linha {i}: {e}")
    return {"ok": True, "total_rows": len(rows), **counters, "errors": errors[:10]}


# ---------------------------------------------------------------------------
# Upload e exibição de fotos de produtos (Object Storage)
# ---------------------------------------------------------------------------
@api_router.post("/upload")
async def upload_image(file: UploadFile = File(...), _: None = Depends(require_admin)):
    data = await file.read()
    if len(data) > 6 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Imagem muito grande (máx. 6 MB)")
    if not (file.content_type or "").startswith("image/"):
        raise HTTPException(status_code=400, detail="Envie um arquivo de imagem")
    ext = (file.filename or "foto.jpg").rsplit(".", 1)[-1].lower()
    if not re.match(r"^[a-z0-9]{2,5}$", ext):
        ext = "jpg"
    path = f"{APP_NAME}/products/{uuid.uuid4().hex}.{ext}"
    try:
        result = await run_in_threadpool(put_object, path, data, file.content_type or "image/jpeg")
    except requests.HTTPError as e:
        code = e.response.status_code if e.response is not None else 500
        msg = "Créditos de armazenamento esgotados" if code == 402 else "Falha ao enviar a foto"
        raise HTTPException(status_code=502, detail=msg)
    stored = result.get("path", path)
    await db.uploads.insert_one({
        "path": stored, "original_name": file.filename, "content_type": file.content_type,
        "size": len(data), "created_at": datetime.now(timezone.utc),
    })
    return {"url": f"/files/{stored}", "path": stored}


@api_router.get("/files/{file_path:path}")
async def serve_file(file_path: str):
    # Catálogo público: só serve fotos de produtos registradas no banco
    found = await db.uploads.find_one({"path": file_path})
    if not found:
        raise HTTPException(status_code=404, detail="Arquivo não encontrado")
    try:
        content, ctype = await run_in_threadpool(get_object, file_path)
    except Exception:
        raise HTTPException(status_code=404, detail="Arquivo não encontrado")
    return Response(content, media_type=ctype, headers={"Cache-Control": "public, max-age=86400"})


app.include_router(api_router)
app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)
